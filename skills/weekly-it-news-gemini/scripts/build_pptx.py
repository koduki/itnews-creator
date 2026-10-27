#!/usr/bin/env python3
"""Build an editable Weekly IT News PPTX from an approved plot and canonical template.

Run in Gemini Tasks' Python-capable environment (python-pptx, Pillow); optional
visual QA requires LibreOffice and pypdfium2. Input images are local files
obtained from Drive/GAS or DOCX extraction, not browser screenshots generated
by the script. The canonical template is not bundled in the skill ZIP.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

EXPECTED_TEMPLATE_SHA256 = 'f871047b53f839a47bb77732770f7648a4e904bc64d696e04e85d2e639778f05'
LAYOUT_TO_TEMPLATE_INDEX = {
    'COVER': 0, 'AGENDA': 1, 'NEWS_TITLE': 2, 'NEWS': 3,
    'SOURCE_IMAGE': 4, 'TECH_DIAGRAM': 5, 'COMPARE': 6, 'TWO_SOURCES': 7,
}

def verify_template(template_path: Path, skip_hash_check: bool = False) -> str:
    """Check the pinned canonical 8-slide presentation and source avatar."""
    if not template_path.is_file():
        raise FileNotFoundError(f'Template PPTX not found: {template_path}')
    digest = hashlib.sha256(template_path.read_bytes()).hexdigest()
    if not skip_hash_check and digest != EXPECTED_TEMPLATE_SHA256:
        raise ValueError(f'Template SHA-256 mismatch: expected {EXPECTED_TEMPLATE_SHA256}, got {digest}')
    with zipfile.ZipFile(template_path) as zf:
        if zf.testzip() is not None:
            raise ValueError('Corrupt template ZIP')
        if 'ppt/media/image1.jpeg' not in zf.namelist():
            raise ValueError('Template is missing original cover avatar ppt/media/image1.jpeg')
    prs = Presentation(str(template_path))
    if len(prs.slides) != 8:
        raise ValueError(f'Expected 8 template slides, got {len(prs.slides)}')
    for name in ('Big Number', 'News Label', 'News Title'):
        if not any(sh.name == name for sh in prs.slides[2].shapes):
            raise ValueError(f'NEWS_TITLE template slide missing {name}')
    return digest


def extract_docx_assets(docx_paths: list[Path], out_dir: Path) -> list[dict]:
    """Extract all embedded PNG/JPEG and document-level hyperlinks from DOCX."""
    out_dir.mkdir(parents=True, exist_ok=True)
    extracted = []
    for docx_path in docx_paths:
        prefix = docx_path.stem.split('_')[0]
        with zipfile.ZipFile(docx_path) as zf:
            rels_urls = {}
            if 'word/_rels/document.xml.rels' in zf.namelist():
                root = ET.fromstring(zf.read('word/_rels/document.xml.rels'))
                for rel in root:
                    if 'hyperlink' in rel.attrib.get('Type', ''):
                        rels_urls[rel.attrib.get('Id', '')] = rel.attrib.get('Target', '')
            for name in sorted(zf.namelist()):
                if name.startswith('word/media/') and Path(name).suffix.lower() in {'.png', '.jpg', '.jpeg'}:
                    out_file = out_dir / f'{prefix}_{Path(name).name}'
                    out_file.write_bytes(zf.read(name))
                    extracted.append({'source_docx': str(docx_path), 'media_name': name,
                                      'local_file': str(out_file), 'hyperlinks_in_doc': list(rels_urls.values())})
    return extracted


def set_cjk_font(run_or_para, font_name: str = 'Yu Gothic') -> None:
    """Only set the East-Asian fallback. Never overwrite template Latin/CS fonts or colors."""
    el = run_or_para._r.get_or_add_rPr() if hasattr(run_or_para, '_r') else run_or_para._p.get_or_add_pPr()
    if el.tag.endswith('pPr'):
        def_rpr = el.find(qn('a:defRPr'))
        if def_rpr is None:
            def_rpr = parse_xml(f'<a:defRPr {nsdecls("a")}/>')
            el.append(def_rpr)
        el = def_rpr
    child = el.find(qn('a:ea'))
    if child is None:
        child = parse_xml(f'<a:ea {nsdecls("a")} typeface="{font_name}"/>')
        el.append(child)
    # Existing a:ea and a:latin (e.g. masthead fonts) are deliberately left intact.

def enforce_slide_cjk_fonts(slide, font_name: str = 'Yu Gothic') -> None:
    """Preserve typography defined in the PPTX template, filling only missing CJK fallback."""
    for shape in slide.shapes:
        text_frames = [shape.text_frame] if shape.has_text_frame else []
        if shape.has_table:
            text_frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
        for tf in text_frames:
            for para in tf.paragraphs:
                set_cjk_font(para, font_name)
                for run in para.runs:
                    set_cjk_font(run, font_name)

def clone_template_slide(prs: Presentation, template_idx: int):
    """Copy native shapes and remap related media IDs, including the cover avatar."""
    src = prs.slides[template_idx]
    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    new = prs.slides.add_slide(blank_layout)
    for shape in list(new.shapes):
        elem = shape._element
        elem.getparent().remove(elem)
    rmap = {}
    for rel in src.part.rels.values():
        if 'notesSlide' in rel.reltype or 'slideLayout' in rel.reltype:
            continue
        if rel.is_external:
            new_id = new.part.rels.get_or_add_ext_rel(rel.reltype, rel._target)
        else:
            new_id = new.part.rels.get_or_add(rel.reltype, rel._target)
        rmap[rel.rId] = new_id
    for shape in src.shapes:
        elem = copy.deepcopy(shape._element)
        for blip in elem.iter(qn('a:blip')):
            old = blip.get(qn('r:embed'))
            if old in rmap:
                blip.set(qn('r:embed'), rmap[old])
        new.shapes._spTree.insert_element_before(elem, 'p:extLst')
    return new


def delete_first_n_slides(prs: Presentation, n: int = 8) -> None:
    ids = prs.slides._sldIdLst
    for el in list(ids)[:n]:
        prs.part.drop_rel(el.rId)
        ids.remove(el)


def find_shape(slide, name: str):
    return next((x for x in slide.shapes if x.name == name), None)


def remove_shape(slide, name: str) -> None:
    sh = find_shape(slide, name)
    if sh is not None:
        elem = sh._element
        elem.getparent().remove(elem)


def compute_title_pt(title: str) -> float | None:
    """Keep the template's native 30pt when possible; only shrink unusually long headings."""
    n = len(title)
    return 18.0 if n >= 45 else 20.0 if n >= 39 else 23.0 if n >= 33 else 26.0 if n >= 23 else None

def _replace_paragraph_text(para, text: str, pt: float | None = None,
                            bold: bool | None = None) -> None:
    """Retain paragraph and run styling (a:defRPr / colors / fonts) from the template."""
    runs = list(para.runs)
    if runs:
        main = runs[0]
        main.text = str(text)
        for other in runs[1:]:
            other._r.getparent().remove(other._r)
    else:
        main = para.add_run()
        main.text = str(text)
    if pt is not None:
        main.font.size = Pt(pt)
    if bold is not None:
        main.font.bold = bold
    set_cjk_font(para)
    set_cjk_font(main)


def set_single_text(shape, text: str, pt: float | None = None,
                    bold: bool | None = None, align: PP_ALIGN | None = None) -> None:
    """Replace text but preserve the template's native fill, font colors and placement."""
    if shape is None or not hasattr(shape, 'text_frame'):
        return
    tf = shape.text_frame
    # Keep its first <a:p> and pPr rather than tf.clear(), which discards
    # the native blue <a:defRPr> style used throughout the eight-slide template.
    while len(tf.paragraphs) > 1:
        node = tf.paragraphs[-1]._p
        node.getparent().remove(node)
    para = tf.paragraphs[0]
    if align is not None:
        para.alignment = align
    _replace_paragraph_text(para, str(text), pt=pt, bold=bold)


def set_template_lines(shape, lines, *, font_size: float | None = None,
                       space_after: float | None = None, bullets: bool = False) -> None:
    """Populate a template text frame; clone styled paragraphs instead of recreating them."""
    if shape is None:
        return
    tf = shape.text_frame
    original = [copy.deepcopy(para._p) for para in tf.paragraphs]
    if not original:
        raise ValueError(f'Expected a styled text paragraph in {shape.name}')
    content = [str(x) for x in lines] or ['']
    while len(tf.paragraphs) < len(content):
        proto = original[min(len(tf.paragraphs), len(original)-1)]
        tf._txBody.append(copy.deepcopy(proto))
    while len(tf.paragraphs) > len(content):
        node = tf.paragraphs[-1]._p
        node.getparent().remove(node)
    tf.word_wrap = True
    for i, line in enumerate(content):
        if bullets and line and not line.startswith('・'):
            line = '・' + line
        para = tf.paragraphs[i]
        if space_after is not None:
            para.space_after = Pt(space_after)
        _replace_paragraph_text(para, line, pt=font_size)

def set_header_and_footer(slide, data: dict, edition_date: str) -> None:
    """The template is the design authority. Update its title, not its colors/geometry."""
    sh = find_shape(slide, 'TextBox 1')
    if sh:
        title = str(data.get('title', ''))
        set_single_text(sh, title, compute_title_pt(title))
    # TextBox 3 is optional in user-supplied templates. Never create a footer
    # or move a template element to enforce legacy coordinates.
    sh = find_shape(slide, 'TextBox 3')
    if sh:
        footer = f"Weekly IT News！ {edition_date}  •  {data.get('slide_id','').upper()}"
        set_single_text(sh, footer)

def populate_news_title(slide, data: dict, number: int) -> None:
    """Maintain the NEWS_TITLE shapes and blue/navy typography from the actual template."""
    label = find_shape(slide, 'News Label')
    big = find_shape(slide, 'Big Number')
    title = find_shape(slide, 'News Title')
    if not all((label, big, title)):
        raise ValueError('NEWS_TITLE native shapes missing')
    nn = f'{number:02d}'
    set_single_text(label, f'NEWS {nn}')
    set_single_text(big, nn)
    headline = str(data.get('title') or '（タイトル未設定）')
    size = None if len(headline) <= 18 else 41 if len(headline) <= 26 else 34 if len(headline) <= 38 else 28
    set_single_text(title, headline, size)
    title.text_frame.word_wrap = True

def expanded_story_slides(plot: dict) -> list[dict]:
    """Generate exactly one story divider before each news block.

    Explicit NEWS_TITLE slides take precedence; older approved plots without
    dividers are upgraded deterministically. The plot remains immutable.
    """
    source = plot.get('slides', [])
    story_titles = plot.get('story_titles') or {}
    if not isinstance(story_titles, dict):
        raise ValueError('story_titles must be a mapping of story_id to title')
    output = []
    last_story = None
    news_number = 0
    for spec in source:
        layout = spec.get('layout_id', 'NEWS')
        story = str(spec.get('story_id') or '').strip()
        regular_story = story not in ('', 'front') and layout not in ('COVER', 'AGENDA')
        if regular_story and story != last_story:
            news_number += 1
            if layout != 'NEWS_TITLE':
                heading = story_titles.get(story) or spec.get('story_title') or spec.get('title') or story
                output.append({'slide_id':f'{spec.get("slide_id", story)}-title',
                               'story_id':story, 'layout_id':'NEWS_TITLE',
                               'title':heading, 'on_slide':[],
                               'speaker_notes':spec.get('story_intro_notes', ''),
                               'asset_ids':[], '_generated_title':True,
                               '_story_number':news_number})
        new_spec = dict(spec)
        if layout == 'NEWS_TITLE':
            new_spec['_story_number'] = news_number
        output.append(new_spec)
        if regular_story:
            last_story = story
    return output


def embed_picture_in_box(slide, img_path: Path, box_left: float, box_top: float,
                         box_w: float, box_h: float) -> None:
    with Image.open(img_path) as img:
        iw, ih = img.size
        img.verify()
    if iw <= 0 or ih <= 0:
        raise ValueError(f'Invalid image dimensions: {img_path}')
    aspect = iw / ih
    if aspect > box_w / box_h:
        w, h = box_w, box_w / aspect
    else:
        w, h = box_h * aspect, box_h
    slide.shapes.add_picture(str(img_path), Inches(box_left + (box_w-w)/2),
                             Inches(box_top+(box_h-h)/2), width=Inches(w), height=Inches(h))


def resolve_asset_file(asset_id: str, assets_map: dict[str, dict], assets_dir: Path | None) -> Path | None:
    asset = assets_map.get(asset_id, {})
    rel = asset.get('file') or asset.get('local_file')
    candidates: list[Path] = []
    if rel:
        candidates.append(Path(rel))
        if assets_dir:
            candidates.extend([assets_dir / rel, assets_dir / Path(rel).name])
    if assets_dir:
        for ext in ('.png', '.jpg', '.jpeg'):
            candidates.extend([assets_dir / f'{asset_id}{ext}', assets_dir / 'downloads' / f'{asset_id}{ext}'])
    return next((c for c in candidates if c.is_file()), None)



# Metadata is an audit trail, not automatic image recognition. A PNG of a
# Cloudflare challenge is a valid PNG but NOT valid evidence for a news article.
BAD_VISUAL_FLAGS = {
    'bot_challenge', 'cookie_overlay', 'ad_dominant', 'wrong_article',
    'unreadable', 'missing_article', 'needs_crop', 'unrelated',
}


def preflight_plot(plot: dict, assets_dir: Path | None = None,
                   strict_assets: bool = False) -> dict:
    """Check storyboard integrity and asset state BEFORE creating a PPTX.

    In strict mode, an image can be embedded only after a visual review of its
    actual pixels. No algorithm here pretends to identify cookie banners or
    Cloudflare challenges from pixels; the reviewer records quality_flags.
    """
    errors: list[str] = []
    warnings: list[str] = []
    inspected: list[dict] = []
    specs = plot.get('slides')
    if not isinstance(specs, list) or not specs:
        errors.append('plot.slides must contain at least one slide')
        specs = []
    assets_list = plot.get('assets', [])
    if not isinstance(assets_list, list):
        errors.append('plot.assets must be a list')
        assets_list = []
    amap: dict[str, dict] = {}
    for asset in assets_list:
        if not isinstance(asset, dict):
            errors.append('Invalid asset record (expected object)')
            continue
        aid = asset.get('id') or asset.get('asset_id')
        if not isinstance(aid, str) or not aid.strip():
            errors.append('Asset missing id/asset_id')
        elif aid in amap:
            errors.append(f'Duplicate asset_id: {aid}')
        else:
            amap[aid] = asset

    seen_slides: set[str] = set()
    closed_stories: set[str] = set()
    active_story = ''
    for index, spec in enumerate(specs, 1):
        if not isinstance(spec, dict):
            errors.append(f'Slide {index} must be a JSON object')
            continue
        sid = spec.get('slide_id') or f'#{index}'
        if sid in seen_slides:
            errors.append(f'Duplicate slide_id: {sid}')
        seen_slides.add(sid)
        layout = spec.get('layout_id', 'NEWS')
        if layout not in LAYOUT_TO_TEMPLATE_INDEX:
            errors.append(f'{sid}: unsupported layout_id {layout}')
        story = str(spec.get('story_id', '')).strip()
        if story != active_story:
            if active_story and active_story != 'front':
                closed_stories.add(active_story)
            if story in closed_stories and story != 'front':
                errors.append(f'{sid}: story_id {story} reopens after another topic')
            active_story = story
        if layout == 'NEWS_TITLE':
            if story in ('', 'front'):
                errors.append(f'{sid}: NEWS_TITLE requires a real story_id')
            elif index > 1 and specs[index-2].get('story_id') == story:
                errors.append(f'{sid}: NEWS_TITLE must be the first slide of its story')
        if index == 1 and layout != 'COVER':
            warnings.append('First slide is not COVER')
        if index == 2 and layout != 'AGENDA':
            warnings.append('Second slide is not AGENDA')
        ids = spec.get('asset_ids', [])
        if not isinstance(ids, list):
            errors.append(f'{sid}: asset_ids must be a list')
            continue
        expected = {'SOURCE_IMAGE': 1, 'TWO_SOURCES': 2}.get(layout, 0)
        if expected and len(ids) != expected:
            msg = f'{sid}: {layout} requires {expected} asset_id(s), got {len(ids)}'
            (errors if strict_assets else warnings).append(msg)
        if layout == 'COMPARE':
            headers, rows, _, _ = parse_compare_rows(spec)
            if not headers or not rows or any(len(row) != len(headers) for row in rows):
                errors.append(f'{sid}: invalid COMPARE table headers/rows')
            elif len(rows) >= 5 or any(len(str(v)) > 150 for row in rows for v in row):
                warnings.append(f'{sid}: dense COMPARE table requires projection-scale review')
        if layout == 'TECH_DIAGRAM' and sum(len(str(v)) for v in spec.get('on_slide', [])) > 460:
            warnings.append(f'{sid}: dense TECH_DIAGRAM may overflow its 3 native boxes')
        for aid in ids[:expected or len(ids)]:
            if not isinstance(aid, str) or not aid:
                errors.append(f'{sid}: empty/invalid asset_id')
                continue
            asset = amap.get(aid)
            if asset is None:
                (errors if strict_assets else warnings).append(f'{sid}: undeclared asset {aid}')
                continue
            path = resolve_asset_file(aid, amap, assets_dir)
            if path is None:
                (errors if strict_assets else warnings).append(f'{sid}: missing file for asset {aid}')
                continue
            try:
                with Image.open(path) as img:
                    dims = img.size
                    img.verify()
                if dims[0] < 1 or dims[1] < 1:
                    raise ValueError('zero dimension')
            except Exception as exc:
                errors.append(f'{sid}: corrupt/not an image for asset {aid}: {exc}')
                continue
            if path.suffix.lower() not in ('.png', '.jpeg', '.jpg'):
                errors.append(f'{sid}: unsupported image format for {aid}: {path.suffix}')
            sha = hashlib.sha256(path.read_bytes()).hexdigest()
            if asset.get('sha256') and asset['sha256'].lower() != sha:
                errors.append(f'{sid}: SHA-256 mismatch for asset {aid}')
            raw_flags = asset.get('quality_flags') or []
            # Accept both a single string and a list; never silently drop a
            # quality flag because a serializer chose a different JSON shape.
            if isinstance(raw_flags, str):
                raw_flags = [raw_flags]
            if not isinstance(raw_flags, list) or any(not isinstance(x, str) for x in raw_flags):
                errors.append(f'{sid}: invalid quality_flags for {aid}')
                raw_flags = []
            bad = set(raw_flags) & BAD_VISUAL_FLAGS
            if bad:
                (errors if strict_assets else warnings).append(
                    f'{sid}: rejected screenshot {aid}, quality_flags={sorted(bad)}')
            if asset.get('status') in ('pending', 'unavailable', 'rejected', 'failed'):
                (errors if strict_assets else warnings).append(
                    f'{sid}: asset {aid} status={asset["status"]}')
            if strict_assets:
                if asset.get('review_status') != 'approved':
                    errors.append(f'{sid}: asset {aid} has no explicit visual approval')
                if not str(asset.get('review_notes') or '').strip():
                    errors.append(f'{sid}: asset {aid} missing visual review_notes')
                if not (asset.get('source_url') or asset.get('source_docx')):
                    errors.append(f'{sid}: asset {aid} missing source_url/source_docx')
            if min(dims) < 500:
                warnings.append(f'{sid}: low resolution image {aid}: {dims[0]}x{dims[1]}')
            inspected.append({'slide_id': sid, 'asset_id': aid, 'file': str(path),
                              'sha256': sha, 'dimensions': list(dims),
                              'review_status': asset.get('review_status', 'unreviewed')})
    return {'passed': not errors, 'strict_assets': strict_assets,
            'errors': errors, 'warnings': list(dict.fromkeys(warnings)),
            'inspected_assets': inspected,
            'slides_count': len(specs), 'assets_inspected': len(inspected),
            'visual_qa': 'must manually inspect all slide previews; not auto-certified'}


def merge_asset_manifest(plot: dict, manifest_path: Path) -> dict:
    """Merge the Drive/GAS inspection manifest with plot assets by immutable id.

    The manifest supplies review status, final downloaded file and hash. When
    both sides name a source URL they must agree; a replacement article must be
    explicitly approved in plot.json rather than silently substituting it.
    """
    raw = json.loads(manifest_path.read_text(encoding='utf-8'))
    rows = raw if isinstance(raw, list) else raw.get('assets', [])
    if not isinstance(rows, list):
        raise ValueError('assets-manifest.json must contain an assets list')
    clone = copy.deepcopy(plot)
    by_id = {(a.get('id') or a.get('asset_id')): a for a in clone.get('assets', [])}
    for asset in rows:
        aid = asset.get('id') or asset.get('asset_id')
        if not aid:
            raise ValueError('assets-manifest.json contains entry without asset_id')
        original = by_id.get(aid, {})
        if original.get('source_url') and asset.get('source_url') and original['source_url'] != asset['source_url']:
            raise ValueError(f'Asset {aid}: source_url differs from approved plot (reapproval required)')
        by_id[aid] = {**original, **asset, 'id': aid}
    clone['assets'] = list(by_id.values())
    return clone


def write_qa_report(report: dict, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def strip_prefix(text: str, prefixes: tuple[str, ...]) -> str:
    return next((text[len(p):].strip() for p in prefixes if text.startswith(p)), text)


def populate_cover(slide, data: dict, edition_date: str) -> None:
    y, m, d = edition_date.split('-')
    # Cover exception: preserve its authentic red/black masthead and avatar.
    masthead = find_shape(slide, 'TextBox 9')
    if masthead is None or 'Weekly' not in masthead.text:
        raise ValueError('Template cover masthead is missing')
    # Adjust date positions ONLY to avoid collision with the original masthead.
    for name, top in (('Rectangle 6',3.13),('Rectangle 5',3.13),
                      ('Rectangle 1',3.83),('TextBox 4',4.48)):
        sh = find_shape(slide, name)
        if sh:
            sh.top = Inches(top)
    set_single_text(find_shape(slide, 'Rectangle 5'), f'{y}年', 21)
    set_single_text(find_shape(slide, 'Rectangle 1'), f'{int(m):02d}月', 46)
    set_single_text(find_shape(slide, 'TextBox 4'), f'{int(d):02d}', 108)
    set_template_lines(find_shape(slide, 'TextBox 13'), data.get('on_slide', []),
                       font_size=19, space_after=10, bullets=True)

def populate_agenda(slide, data: dict, edition_date: str) -> None:
    set_header_and_footer(slide, data, edition_date)
    set_template_lines(find_shape(slide, 'TextBox 4'), data.get('on_slide', []),
                       font_size=24, space_after=22, bullets=True)

def populate_news(slide, data: dict, edition_date: str) -> None:
    set_header_and_footer(slide, data, edition_date)
    lines = list(data.get('on_slide', []))
    intro, bullets, source = '', [], ''
    for line in lines:
        if line.startswith('導入：'):
            intro = strip_prefix(line, ('導入：',))
        elif line.startswith('出典：'):
            source = line
        else:
            bullets.append(line)
    if not intro and bullets:
        intro = strip_prefix(bullets.pop(0), ('導入：', '・'))
    if not source and data.get('source_urls'):
        source = '出典：' + ' / '.join(data['source_urls'][:3])
    set_single_text(find_shape(slide, 'TextBox 4'), intro, 19)
    set_template_lines(find_shape(slide, 'TextBox 5'), bullets, font_size=17.5,
                       space_after=12, bullets=True)
    set_single_text(find_shape(slide, 'TextBox 6'), source, 11)

def _embed_in_template_frame(slide, path: Path, rectangle_name: str):
    """Use the native frame shape's geometry; avoid hard-coded coordinates."""
    frame = find_shape(slide, rectangle_name)
    if frame is None:
        raise ValueError(f'Missing image frame: {rectangle_name}')
    pad = Inches(0.08)
    embed_picture_in_box(slide, path,
        (frame.left+pad)/Inches(1), (frame.top+pad)/Inches(1),
        (frame.width-2*pad)/Inches(1), (frame.height-2*pad)/Inches(1))


def populate_source_image(slide, data: dict, edition_date: str,
                          assets_map: dict[str, dict], assets_dir: Path | None) -> None:
    set_header_and_footer(slide, data, edition_date)
    lines = list(data.get('on_slide', []))
    desc = lines[0] if lines else ''
    point = lines[1] if len(lines) >= 2 else desc
    source = next((line for line in lines if line.startswith('出典：')), '')
    if not source and data.get('source_urls'):
        source = '出典：' + ' / '.join(data['source_urls'][:2])
    set_single_text(find_shape(slide, 'TextBox 6'), point, 15.5)
    set_single_text(find_shape(slide, 'TextBox 7'), source, 11)
    asset_id = next(iter(data.get('asset_ids', [])), '')
    path = resolve_asset_file(asset_id, assets_map, assets_dir) if asset_id else None
    if path:
        remove_shape(slide, 'TextBox 5')
        _embed_in_template_frame(slide, path, 'Rectangle 4')
    else:
        urls = data.get('source_urls', [])
        url = urls[0] if urls else 'URL未設定'
        placeholder = (f'【記事スクリーンショット枠：{asset_id or "pending"}】\n'
                       f'{desc}\nURL: {url}\n※画像ファイル未取得（差替用）')
        set_single_text(find_shape(slide, 'TextBox 5'), placeholder, 15)

def populate_tech_diagram(slide, data: dict, edition_date: str) -> None:
    set_header_and_footer(slide, data, edition_date)
    lines = list(data.get('on_slide', []))
    lead = strip_prefix(lines[0], ('説明文：',)) if lines else ''
    blocks = [x for x in lines[1:] if x.startswith('ブロック')]
    if len(blocks) < 3 and len(lines) >= 4:
        blocks = lines[1:4]
    bottom = lines[-1] if len(lines) >= 5 else ''
    set_single_text(find_shape(slide, 'TextBox 4'), lead, 17)
    frames = ('Rectangle 5', 'Rectangle 7', 'Rectangle 9')
    boxes = ('TextBox 6', 'TextBox 8', 'TextBox 10')
    for i in range(3):
        rect = find_shape(slide, frames[i])
        txt = find_shape(slide, boxes[i])
        if txt is None or rect is None:
            raise ValueError(f'Missing TECH_DIAGRAM native template shapes: {frames[i]} / {boxes[i]}')
        # Expand the text within the actual native box instead of moving/repainting it.
        inset = Inches(0.14)
        txt.left, txt.top = rect.left + inset, rect.top + inset
        txt.width, txt.height = rect.width - 2*inset, rect.height - 2*inset
        raw = blocks[i] if i < len(blocks) else ''
        heading, body = raw, ''
        if ' — ' in raw:
            heading, body = raw.split(' — ', 1)
            heading = re.sub(r'^ブロック\d+[：:]\s*', '', heading)
        elif '）：' in raw:
            m = re.match(r'^ブロック\d+（(.+?)）：(.*)$', raw)
            if m:
                heading, body = m.group(1), m.group(2)
        title = f'{"①②③"[i]} {heading}' if not heading.startswith(('①','②','③','層')) else heading
        segments = [x.strip() for x in re.split(r'(?<=[。】])\s*|(?<=、)(?=①|②|③|④)', body) if x.strip()]
        set_template_lines(txt, [title] + [s if s.startswith(('【','・')) else '・'+s for s in segments],
                           font_size=14.5 if len(raw) > 90 else 16, space_after=5)
    # Native template arrow styling (blue) is retained.
    set_single_text(find_shape(slide, 'TextBox 11'), '→')
    set_single_text(find_shape(slide, 'TextBox 12'), '→')
    set_single_text(find_shape(slide, 'TextBox 13'), bottom, 12)

def parse_compare_rows(data: dict) -> tuple[list[str], list[list[str]], str, str]:
    lines = list(data.get('on_slide', []))
    lead = source = ''
    data_lines = []
    for line in lines:
        if line.startswith(('注記：', '前提注記：')):
            lead = strip_prefix(line, ('注記：', '前提注記：'))
        elif line.startswith('出典：'):
            source = line
        else:
            data_lines.append(line)
    if isinstance(data.get('table'), dict):
        t = data['table']
        return t.get('headers', []), t.get('rows', []), lead, source
    if any('｜' in x for x in data_lines):
        rows = []
        for line in data_lines:
            if '｜' in line:
                rows.append([c.strip() for c in re.sub(r'^行\d+[：:]\s*', '', line).split('｜')])
        headers = ['プロダクト名', '発表日 / 利用開始日', '価格・プラン', '主な特徴・セキュリティ統制']
        if rows and len(rows[0]) != 4:
            headers = [f'項目{i+1}' for i in range(len(rows[0]))]
        return headers, rows, lead, source
    headers = ['比較観点', 'Log4Shell（2021年12月）', 'PHP2Shell（2026年10月）']
    rows = []
    for line in data_lines:
        m = re.match(r'^比較観点\d+（(.+?)）：(.*)$', line)
        if m:
            aspect, rest = m.group(1), m.group(2)
            if '／' in rest:
                lhs, rhs = [p.strip() for p in rest.split('／', 1)]
                rows.append([aspect, re.sub(r'^Log4Shell＝', '', lhs), re.sub(r'^PHP2Shell＝', '', rhs)])
            else:
                rows.append([aspect, rest, rest])
    return headers, rows, lead, source


def populate_compare(slide, data: dict, edition_date: str) -> None:
    """Resize/copy the template's editable blue table; do not generate a red replacement."""
    set_header_and_footer(slide, data, edition_date)
    headers, rows, lead, source = parse_compare_rows(data)
    set_single_text(find_shape(slide, 'TextBox 4'), lead, 16)
    set_single_text(find_shape(slide, 'TextBox 6'), source, 11.5)
    tables = [sh for sh in slide.shapes if sh.has_table]
    if len(tables) != 1:
        raise ValueError('COMPARE requires exactly one styled template table')
    if not headers or not rows:
        return
    sh = tables[0]
    table = sh.table
    xml_table = table._tbl
    grid = xml_table.tblGrid
    original_grid = list(grid)
    row_elems = list(xml_table.tr_lst)
    head_proto = copy.deepcopy(row_elems[0])
    body_proto = copy.deepcopy(row_elems[1] if len(row_elems) > 1 else row_elems[0])
    ncols = len(headers)
    nrows = len(rows) + 1
    # Grow columns with cloned native cell properties (pale blue + blue text).
    while len(grid) < ncols:
        grid.append(copy.deepcopy(original_grid[-1]))
    while len(grid) > ncols:
        grid.remove(grid[-1])
    # Original 3-col table uses its own proportions; 4-col tables give more
    # space to the final descriptive column while retaining all native colors.
    weights = ([float(x.w) for x in original_grid[:ncols]] if ncols == len(original_grid)
               else [0.20,0.20,0.20,0.40] if ncols == 4
               else [1]*ncols)
    total = sh.width
    for col, weight in zip(table.columns, weights):
        col.width = int(total * weight / sum(weights))
    while len(xml_table.tr_lst) < nrows:
        xml_table.append(copy.deepcopy(body_proto))
    while len(xml_table.tr_lst) > nrows:
        xml_table.remove(xml_table.tr_lst[-1])
    for idx, tr in enumerate(xml_table.tr_lst):
        proto = head_proto if idx == 0 else body_proto
        cells = tr.tc_lst
        while len(cells) < ncols:
            tr.append(copy.deepcopy(proto.tc_lst[-1]))
            cells = tr.tc_lst
        while len(cells) > ncols:
            tr.remove(cells[-1])
            cells = tr.tc_lst
    # Fit the table between its native top and source note, with no overlap.
    note = find_shape(slide, 'TextBox 6')
    available = int((note.top - Inches(0.18) - sh.top) if note is not None else sh.height)
    target_height = min(available, max(sh.height, Inches(0.63*nrows)))
    height_per_row = int(target_height / nrows)
    for row in table.rows:
        row.height = height_per_row
    for j, val in enumerate(headers):
        set_single_text(table.cell(0,j), val, 14 if nrows >= 6 else None)
    for i,row_vals in enumerate(rows,1):
        for j in range(ncols):
            set_single_text(table.cell(i,j), row_vals[j] if j<len(row_vals) else '',
                            12.5 if nrows >= 6 else None)

def populate_two_sources(slide, data: dict, edition_date: str,
                         assets_map: dict[str, dict], assets_dir: Path | None) -> None:
    set_header_and_footer(slide, data, edition_date)
    lines = data.get('on_slide', [])
    set_single_text(find_shape(slide, 'TextBox 8'), lines[0] if lines else '', 13.5)
    set_single_text(find_shape(slide, 'TextBox 9'), lines[1] if len(lines)>1 else '', 13.5)
    ids = data.get('asset_ids', [])
    for idx, ph, frame in [(0,'TextBox 6','Rectangle 4'),(1,'TextBox 7','Rectangle 5')]:
        aid = ids[idx] if idx < len(ids) else ''
        path = resolve_asset_file(aid, assets_map, assets_dir) if aid else None
        if path:
            remove_shape(slide, ph)
            _embed_in_template_frame(slide, path, frame)
        else:
            set_single_text(find_shape(slide, ph), f'【画像未取得：{aid or "pending"}】', 14)

def set_speaker_notes(slide, notes_text: str) -> None:
    if notes_text:
        slide.notes_slide.notes_text_frame.text = notes_text


def build_deck(template_path: Path, plot_path: Path, out_path: Path,
               assets_dir: Path | None = None, skip_hash_check: bool = False,
               strict_assets: bool = False, qa_report_path: Path | None = None,
               assets_manifest_path: Path | None = None) -> dict:
    template_sha = verify_template(template_path, skip_hash_check)
    plot = json.loads(plot_path.read_text(encoding='utf-8'))
    if plot.get('status') != 'approved':
        raise ValueError(f"Plot status must be 'approved' before generating PPTX (got {plot.get('status')!r})")
    if assets_manifest_path is not None:
        plot = merge_asset_manifest(plot, assets_manifest_path)
    report = preflight_plot(plot, assets_dir, strict_assets)
    if qa_report_path:
        write_qa_report(report, qa_report_path)
    if report['errors']:
        raise ValueError('PPTX preflight failed: ' + '; '.join(report['errors']))
    edition_date = plot.get('edition_date', '2026-10-26')
    specs = expanded_story_slides(plot)
    if not specs:
        raise ValueError('plot.slides must contain at least one slide')
    report['output_slides_count'] = len(specs)
    report['generated_news_titles'] = sum(1 for x in specs if x.get('_generated_title'))
    if qa_report_path:
        write_qa_report(report, qa_report_path)
    amap = {(a.get('id') or a.get('asset_id')): a for a in plot.get('assets', []) if a.get('id') or a.get('asset_id')}
    prs = Presentation(str(template_path))
    funcs = {'COVER': populate_cover, 'AGENDA': populate_agenda, 'NEWS': populate_news,
             'TECH_DIAGRAM': populate_tech_diagram, 'COMPARE': populate_compare}
    for spec in specs:
        lay = spec.get('layout_id', 'NEWS')
        if lay not in LAYOUT_TO_TEMPLATE_INDEX:
            raise ValueError(f'Unsupported layout_id: {lay}')
        slide = clone_template_slide(prs, LAYOUT_TO_TEMPLATE_INDEX[lay])
        if lay in funcs:
            funcs[lay](slide, spec, edition_date)
        elif lay == 'NEWS_TITLE':
            populate_news_title(slide, spec, int(spec.get('_story_number', 1)))
        elif lay == 'SOURCE_IMAGE':
            populate_source_image(slide, spec, edition_date, amap, assets_dir)
        elif lay == 'TWO_SOURCES':
            populate_two_sources(slide, spec, edition_date, amap, assets_dir)
        set_speaker_notes(slide, spec.get('speaker_notes', ''))
        if lay not in ('COVER', 'NEWS_TITLE'):
            enforce_slide_cjk_fonts(slide)  # Keep cover branding and blue news-divider typography.
    delete_first_n_slides(prs)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return {'output_pptx': str(out_path), 'slides_count': len(prs.slides),
            'template_sha256': template_sha, 'preflight': report,
            'news_titles_count': sum(1 for x in specs if x['layout_id'] == 'NEWS_TITLE'),
            'slide_mapping': [{'slide_id': x.get('slide_id'), 'story_id': x.get('story_id'),
                               'layout_id': x.get('layout_id')} for x in specs]}


def render_qa_previews(pptx_path: Path, qa_dir: Path, scale: float = 2.0) -> list[Path]:
    """Render via LibreOffice PDF then pypdfium2 PNG for manual visual review."""
    import pypdfium2 as pdfium
    qa_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(['libreoffice', '-env:UserInstallation=file:///tmp/itnews_qa_lo',
                    '--headless', '--convert-to', 'pdf', '--outdir', str(qa_dir), str(pptx_path)],
                   check=True, capture_output=True)
    pdf = pdfium.PdfDocument(str(qa_dir / (pptx_path.stem+'.pdf')))
    rendered = []
    for idx in range(len(pdf)):
        out = qa_dir / f'slide_{idx+1:02d}.png'
        pdf[idx].render(scale=scale).to_pil().save(out)
        rendered.append(out)
    return rendered


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--template', type=Path, required=True)
    parser.add_argument('--plot', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--assets-dir', type=Path)
    parser.add_argument('--render-qa-dir', type=Path)
    parser.add_argument('--skip-hash-check', action='store_true')
    parser.add_argument('--strict-assets', action='store_true',
                        help='Block missing/unreviewed/wrong evidence before generating')
    parser.add_argument('--qa-report', type=Path, help='Write machine-readable preflight JSON')
    parser.add_argument('--assets-manifest', type=Path,
                        help='Drive/GAS assets-manifest.json with visual review metadata')
    args = parser.parse_args()
    result = build_deck(args.template, args.plot, args.out, args.assets_dir,
                        args.skip_hash_check, args.strict_assets, args.qa_report, args.assets_manifest)
    print(f"Generated {result['output_pptx']} ({result['slides_count']} slides)")
    if args.render_qa_dir:
        files = render_qa_previews(args.out, args.render_qa_dir)
        print(f'Rendered {len(files)} QA slide previews in {args.render_qa_dir}')


if __name__ == '__main__':
    main()