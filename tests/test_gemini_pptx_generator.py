"""Regression tests of the Gemini Tasks editable PPTX renderer.

Synthetic test images are NOT presented as documentary evidence.
"""
from __future__ import annotations

import importlib.util
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'assets/templates/weekly-it-news-design-template.pptx'
SCRIPT = ROOT / 'skills/weekly-it-news-gemini/scripts/build_pptx.py'
MODULE = importlib.util.spec_from_file_location('gemini_generator', SCRIPT)
generator = importlib.util.module_from_spec(MODULE)
MODULE.loader.exec_module(generator)


def sample_plot():
    return {
        'edition_date': '2026-10-26', 'status': 'approved',
        'slides': [
            {'slide_id':'s01','story_id':'front','layout_id':'COVER','title':'Weekly IT News！','on_slide':['検証用見出し','二つ目の見出し'],'speaker_notes':'cover note'},
            {'slide_id':'s02','story_id':'front','layout_id':'AGENDA','title':'今週のトピック','on_slide':['A','B'],'speaker_notes':'agenda note'},
            {'slide_id':'s03','story_id':'a','layout_id':'NEWS','title':'ニュース検証','on_slide':['導入：背景','技術的な違い','出典：test'], 'speaker_notes':'news note'},
            {'slide_id':'s04','story_id':'a','layout_id':'SOURCE_IMAGE','title':'証拠画像（テスト）','on_slide':['画像見出し','画像の要点','出典：test'], 'asset_ids':['image1']},
            {'slide_id':'s05','story_id':'a','layout_id':'TECH_DIAGRAM','title':'技術的な模式図','on_slide':['説明文：構成','ブロック1：入力 — テスト。','ブロック2：処理 — 処理。','ブロック3：結果 — 出力。','出典：test']},
            {'slide_id':'s06','story_id':'b','layout_id':'COMPARE','title':'編集可能な比較表','table':{'headers':['観点','A','B'],'rows':[['速度','速い','遅い']]},'on_slide':['注記：同条件','出典：test']},
            {'slide_id':'s07','story_id':'b','layout_id':'TWO_SOURCES','title':'2つの画像比較','on_slide':['左画像','右画像'],'asset_ids':['image1','image2']},
        ],
        'assets': [{'id':'image1','file':'image1.png'}, {'id':'image2','file':'image2.jpg'}],
    }


class GeminiPptxGeneratorTest(unittest.TestCase):
    def test_template_integrity_and_hash_mismatch(self):
        self.assertEqual(generator.verify_template(TEMPLATE),
                         'f871047b53f839a47bb77732770f7648a4e904bc64d696e04e85d2e639778f05')
        with tempfile.TemporaryDirectory() as td:
            other = Path(td) / 'other.pptx'
            other.write_bytes(TEMPLATE.read_bytes() + b'X')
            with self.assertRaisesRegex(ValueError, 'SHA-256 mismatch'):
                generator.verify_template(other)

    def test_generated_pptx_8_layouts_and_auto_story_titles(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            im1 = Image.new('RGB', (400, 220), 'white')
            im2 = Image.new('RGB', (220, 400), 'white')
            im1.save(root/'image1.png')
            im2.save(root/'image2.jpg')
            p = root/'plot.json'
            p.write_text(json.dumps(sample_plot(), ensure_ascii=False), encoding='utf-8')
            out = root/'deck.pptx'
            result = generator.build_deck(TEMPLATE, p, out, assets_dir=root)
            self.assertEqual(result['slides_count'], 9)
            self.assertEqual(result['news_titles_count'], 2)
            self.assertEqual(result['preflight']['generated_news_titles'], 2)
            with zipfile.ZipFile(out) as z:
                self.assertIsNone(z.testzip())
                media = [f for f in z.namelist() if f.startswith('ppt/media/')]
                # cover avatar + genuine TEST image bytes (2 copies of image1, 1 image2)
                self.assertGreaterEqual(len(media), 3)
                self.assertTrue(any(z.read(name) == (root/'image1.png').read_bytes() for name in media))
                self.assertTrue(any(z.read(name) == (root/'image2.jpg').read_bytes() for name in media))
                self.assertTrue(any(z.read(name) == (ROOT/'assets/cover/avatar-20260315.jpg').read_bytes() for name in media))
                slide_xmls = [z.read(name) for name in z.namelist() if name.startswith('ppt/slides/slide') and name.endswith('.xml')]
                self.assertTrue(any(b'Yu Gothic' in xml for xml in slide_xmls))
            deck = Presentation(out)
            self.assertEqual(len(deck.slides), 9)
            self.assertIn('2026年', '\n'.join(sh.text for sh in deck.slides[0].shapes if sh.has_text_frame))
            self.assertEqual(sum(sh.shape_type == 13 for sh in deck.slides[4].shapes), 1)
            self.assertEqual(sum(sh.shape_type == 13 for sh in deck.slides[8].shapes), 2)
            self.assertEqual(sum(sh.has_table for sh in deck.slides[7].shapes), 1)
            self.assertIn('NEWS 01', '\n'.join(sh.text for sh in deck.slides[2].shapes if sh.has_text_frame))
            self.assertIn('ニュース検証', '\n'.join(sh.text for sh in deck.slides[2].shapes if sh.has_text_frame))
            self.assertIn('NEWS 02', '\n'.join(sh.text for sh in deck.slides[6].shapes if sh.has_text_frame))
            self.assertIn('cover note', deck.slides[0].notes_slide.notes_text_frame.text)
            # Layout IDs are fixed, images are in the deck, and seed slides are gone.
            self.assertNotIn('ニュースA：', '\n'.join(sh.text for sh in deck.slides[1].shapes if sh.has_text_frame))

    def test_explicit_news_title_is_not_duplicated(self):
        sample = sample_plot()
        sample['slides'].insert(2, {'slide_id':'s02-title', 'story_id':'a',
                                    'layout_id':'NEWS_TITLE', 'title':'人が決めたニュース名'})
        expanded = generator.expanded_story_slides(sample)
        self.assertEqual(sum(x['layout_id']=='NEWS_TITLE' for x in expanded), 2)
        self.assertEqual(expanded[2]['title'], '人が決めたニュース名')
        self.assertEqual(expanded[2]['_story_number'], 1)
        self.assertFalse(expanded[2].get('_generated_title', False))
        self.assertEqual(expanded[6]['_story_number'], 2)

    def test_story_title_mapping_does_not_replace_slide_headline(self):
        sample = sample_plot()
        sample['story_titles']={'a':'プロットで承認したニュースAの章題',
                                'b':'ニュースBの章題'}
        expanded = generator.expanded_story_slides(sample)
        self.assertEqual(expanded[2]['title'], 'プロットで承認したニュースAの章題')
        self.assertEqual(expanded[3]['title'], 'ニュース検証')
        self.assertEqual(expanded[6]['title'], 'ニュースBの章題')

    def test_template_colors_survive_generation(self):
        """Cover has its own identity; all other slides must inherit blue/navy."""
        from pptx.dml.color import RGBColor
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            p = root / 'plot.json'
            p.write_text(json.dumps(sample_plot(), ensure_ascii=False), encoding='utf-8')
            output = root / 'deck.pptx'
            generator.build_deck(TEMPLATE, p, output, assets_dir=root)
            result = Presentation(output)
            template = Presentation(TEMPLATE)

            def named(slide, name):
                return next(shape for shape in slide.shapes if shape.name == name)

            def inherited_color(shape):
                para = shape.text_frame.paragraphs[0]._p
                defaults = para.findall('.//' + '{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr')
                self.assertTrue(defaults, shape.name)
                colors = defaults[0].findall('.//' + '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
                self.assertTrue(colors, shape.name)
                return colors[0].get('val')

            native_blue = inherited_color(named(template.slides[3], 'TextBox 1'))
            self.assertEqual(native_blue, '0050E6')
            for i in [1, 3, 4, 5, 7, 8]:
                self.assertEqual(inherited_color(named(result.slides[i], 'TextBox 1')), native_blue)

            # Core body text and diagram labels inherit the template's navy.
            self.assertEqual(inherited_color(named(result.slides[3], 'TextBox 4')), '0A1A3A')
            self.assertEqual(inherited_color(named(result.slides[5], 'TextBox 6')), '0A1A3A')
            self.assertEqual(inherited_color(named(result.slides[5], 'TextBox 11')), native_blue)

            # Created COMPARE table inherits the sample table, NOT a new red header.
            table = next(sh.table for sh in result.slides[7].shapes if sh.has_table)
            self.assertEqual(table.cell(0, 0).fill.fore_color.rgb, RGBColor(0xEA, 0xF2, 0xFF))
            header_color = table.cell(0, 0).text_frame.paragraphs[0]._p.findall(
                './/' + '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
            self.assertIn(native_blue, [c.get('val') for c in header_color])

            # Exception: red remains on the COVER only, where the template uses it.
            for index in range(2, len(result.slides) + 1):
                # OOXML slide part names are not renumbered when seed slides
                # are deleted (e.g. the output begins with slide9.xml).
                xml = result.slides[index - 1]._element.xml
                self.assertNotIn('C00000', xml, f'Unexpected legacy red on slide {index}')

    def test_draft_status_blocks_presentation(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'plot.json'
            plot=sample_plot(); plot['status']='draft'
            p.write_text(json.dumps(plot, ensure_ascii=False),encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'approved'):
                generator.build_deck(TEMPLATE,p,Path(td)/'deck.pptx')

    def test_docx_media_and_relationships_extracted(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);doc=root/'report.docx';data=b'\x89PNG\r\n\x1a\nEXAMPLE'
            with zipfile.ZipFile(doc,'w') as z:
                z.writestr('word/media/pic1.png', data)
                z.writestr('word/_rels/document.xml.rels',
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com" TargetMode="External"/></Relationships>')
            items=generator.extract_docx_assets([doc], root/'assets')
            self.assertEqual(len(items),1)
            self.assertIn('https://example.com',items[0]['hyperlinks_in_doc'])
            self.assertEqual(Path(items[0]['local_file']).read_bytes(),data)


if __name__ == '__main__':
    unittest.main()