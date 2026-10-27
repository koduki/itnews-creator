"""Regression: the real eight-slide template, not a legacy red theme, owns presentation styles."""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from pptx import Presentation
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'assets/templates/weekly-it-news-design-template.pptx'
SCRIPT = ROOT / 'skills/weekly-it-news-gemini/scripts/build_pptx.py'
spec = importlib.util.spec_from_file_location('gemini_blue', SCRIPT)
render = importlib.util.module_from_spec(spec)
spec.loader.exec_module(render)


def xml_color(shape):
    """Obtain canonical hex RGB from the paragraph default style, not a new run."""
    from pptx.oxml.ns import qn
    p = shape.text_frame.paragraphs[0]._p
    matches = p.xpath('./a:pPr/a:defRPr/a:solidFill/a:srgbClr')
    return matches[0].get('val') if matches else None


def fill_hex(cell):
    return str(cell.fill.fore_color.rgb)


class TemplateDesignTest(unittest.TestCase):
    def test_all_text_reuses_blue_template_style(self):
        template=Presentation(TEMPLATE)
        src_title = next(x for x in template.slides[3].shapes if x.name == 'TextBox 1')
        self.assertEqual(xml_color(src_title), '0050E6')
        self.assertEqual(fill_hex(next(x for x in template.slides[6].shapes if x.has_table).table.cell(0,0)), 'EAF2FF')
        with tempfile.TemporaryDirectory() as tmp:
            directory=Path(tmp)
            Image.new('RGB',(1280,853),(128,196,235)).save(directory/'image.png')
            slides=[
                {'slide_id':'1','story_id':'front','layout_id':'COVER','title':'Weekly IT News！','on_slide':['A','B']},
                {'slide_id':'2','story_id':'front','layout_id':'AGENDA','title':'今週のトピック','on_slide':['A','B']},
                {'slide_id':'3','story_id':'a','layout_id':'NEWS','title':'ニュース本文タイトル','on_slide':['導入：説明','比較する話題','出典：https://example.test']},
                {'slide_id':'4','story_id':'a','layout_id':'TECH_DIAGRAM','title':'技術解説','on_slide':['説明文：概要','ブロック1：入力 — 説明。','ブロック2：処理 — 説明。','ブロック3：結果 — 説明。','注記：出典']},
                {'slide_id':'5','story_id':'a','layout_id':'COMPARE','title':'比較表','table':{'headers':['項目','A','B','詳細'],'rows':[['性能','速い','遅い','比較説明'],['コスト','高い','低い','確認']]},'on_slide':['注記：同条件','出典：source']},
                {'slide_id':'6','story_id':'a','layout_id':'SOURCE_IMAGE','title':'実画像','on_slide':['概要','注目点','出典：https://example.test'],'asset_ids':['img']},
                {'slide_id':'7','story_id':'a','layout_id':'TWO_SOURCES','title':'二画面','on_slide':['A','B'],'asset_ids':['img','img']},
            ]
            plot={'status':'approved','edition_date':'2026-10-27','slides':slides,'assets':[{'id':'img','file':'image.png','source_url':'https://example.test'}]}
            data=directory/'plot.json';data.write_text(json.dumps(plot,ensure_ascii=False),encoding='utf-8')
            out=directory/'generated.pptx'
            result=render.build_deck(TEMPLATE,data,out,assets_dir=directory)
            self.assertEqual(result['slides_count'],8)  # one generated NEWS_TITLE
            self.assertEqual(result['preflight']['template_style_audit'], 'PASS')
            self.assertEqual(len(result['preflight']['generator_sha256']), 64)
            self.assertEqual(result['preflight']['template_sha256'], render.EXPECTED_TEMPLATE_SHA256)
            prs=Presentation(out)
            self.assertEqual(len(prs.slides),8)
            for index in (1,3,4,5,6,7):
                title=next(x for x in prs.slides[index].shapes if x.name=='TextBox 1')
                self.assertEqual(xml_color(title),'0050E6',f'slide {index+1} must remain blue')
                self.assertNotIn('C00000', title._element.xml)
            self.assertEqual(xml_color(next(x for x in prs.slides[4].shapes if x.name=='TextBox 11')),
                             xml_color(next(x for x in template.slides[5].shapes if x.name=='TextBox 11')))
            table=next(x for x in prs.slides[5].shapes if x.has_table).table
            self.assertEqual(len(table.columns),4)
            self.assertEqual(len(table.rows),3)
            self.assertEqual(fill_hex(table.cell(0,0)),'EAF2FF')
            self.assertEqual(fill_hex(table.cell(0,3)),'EAF2FF')
            self.assertNotIn('C00000',table.cell(0,0)._tc.xml)
            self.assertEqual(len([x for x in prs.slides[6].shapes if x.shape_type==13]),1)
            self.assertEqual(len([x for x in prs.slides[7].shapes if x.shape_type==13]),2)
            # Native header frame geometry must not be replaced with old absolute positions.
            actual=next(x for x in prs.slides[3].shapes if x.name=='TextBox 1')
            self.assertEqual((actual.left,actual.top,actual.width,actual.height),
                             (src_title.left,src_title.top,src_title.width,src_title.height))
            self.assertIn('Weekly',next(x for x in prs.slides[0].shapes if x.name=='TextBox 9').text)

    def test_runtime_audit_rejects_template_style_override(self):
        """The generator must fail even if the offending color exists elsewhere."""
        from pptx.dml.color import RGBColor
        template = Presentation(TEMPLATE)
        generated = render.clone_template_slide(template, 3)
        spec = [{'slide_id': 'color-check', 'layout_id': 'NEWS'}]
        render.audit_template_styles(template, spec)  # untouched template clone

        title = next(sh for sh in generated.shapes if sh.name == 'TextBox 1')
        para = title.text_frame.paragraphs[0]
        run = para.runs[0] if para.runs else para.add_run()
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        with self.assertRaisesRegex(ValueError, 'template style'):
            render.audit_template_styles(template, spec)

        run.font.color.rgb = RGBColor(0x0A, 0x1A, 0x3A)
        # Navy is already used elsewhere on NEWS slides, but the headline is blue.
        with self.assertRaisesRegex(ValueError, 'template style changed in TextBox 1'):
            render.audit_template_styles(template, spec)

    def test_template_contract_prevents_old_red_rules(self):
        contract=(ROOT/'skills/weekly-it-news-gemini/references/plot-and-layout.md').read_text()
        self.assertIn('デザインの唯一の正本',contract)
        self.assertIn('AGENDA以降はすべて青基調',contract)
        self.assertNotIn('17.5pt 赤Bold',contract)
        polish=(ROOT/'skills/itnews-powerpoint-polish/references/visual-repair.md').read_text()
        self.assertIn('青い見出し',polish)


if __name__=='__main__':
    unittest.main()